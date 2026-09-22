import tkinter as tk
from pathlib import Path
from tkinter import messagebox, simpledialog, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

# Filen ligger i mappen "mails" ved siden av dette scriptet
FILE = str(Path(__file__).parent / "mails" / "mails.xlsx")
KLASSE_KOLONNE = 1  # 0-basert plass i tabellen: mail=0, klasse=1, rolle=2
ROLLE_KOLONNE = 2
GYLDIGE_ROLLER = ("student", "teacher")


# ---------------------------------------------------------------------------
# Excel-funksjoner
# ---------------------------------------------------------------------------

def hent_tabell(ark: Worksheet):
    """Returnerer (tabell, min_kol, min_rad, maks_kol, maks_rad)."""
    tabell = list(ark.tables.values())[0]
    return (tabell, *range_boundaries(tabell.ref))


def _sett_ref(tabell, min_kol: int, min_rad: int, maks_kol: int, maks_rad: int) -> None:
    ref = f"{get_column_letter(min_kol)}{min_rad}:{get_column_letter(maks_kol)}{maks_rad}"
    tabell.ref = ref
    if tabell.autoFilter:
        tabell.autoFilter.ref = ref


def _sjekk_rolle(rolle: str) -> None:
    if rolle not in GYLDIGE_ROLLER:
        gyldige = " eller ".join(f"«{r}»" for r in GYLDIGE_ROLLER)
        raise ValueError(f"Rolle må være {gyldige}, ikke «{rolle}».")


def endre_en(arbeidsbok: Workbook, ark: Worksheet, rad: int, kolonne: int, verdi: str) -> None:
    _, min_kol, *_ = hent_tabell(ark)
    if kolonne == min_kol + ROLLE_KOLONNE:
        _sjekk_rolle(verdi)
    # .value = ... i stedet for cell(value=...) slik at tom streng/None faktisk tømmer cellen
    ark.cell(row=rad, column=kolonne).value = verdi or None
    arbeidsbok.save(FILE)


def endre_rad(arbeidsbok: Workbook, ark: Worksheet, rad: int, verdier: list[str]) -> None:
    """Skriver nye verdier til hele raden (én verdi per kolonne i tabellen) og lagrer én gang."""
    _, min_kol, min_rad, maks_kol, maks_rad = hent_tabell(ark)

    if not (min_rad < rad <= maks_rad):
        raise ValueError(f"Rad {rad} er ikke en datarad i tabellen ({min_rad + 1}-{maks_rad}).")
    if len(verdier) != maks_kol - min_kol + 1:
        raise ValueError("Antall verdier stemmer ikke med antall kolonner i tabellen.")
    if not verdier[0]:
        raise ValueError("Første kolonne kan ikke være tom.")
    _sjekk_rolle(verdier[ROLLE_KOLONNE])

    for kol, verdi in enumerate(verdier, start=min_kol):
        ark.cell(row=rad, column=kol).value = verdi or None
    arbeidsbok.save(FILE)


def endre_alle(arbeidsbok: Workbook, ark: Worksheet, klasse: str, nyklasse: str) -> int:
    """Endrer alle rader med klasse -> nyklasse. Returnerer antall endrede rader."""
    _, min_kol, min_rad, maks_kol, maks_rad = hent_tabell(ark)
    kol = min_kol + KLASSE_KOLONNE

    antall = 0
    for r in range(min_rad + 1, maks_rad + 1):
        celle = ark.cell(row=r, column=kol)
        if celle.value is not None and str(celle.value) == klasse:
            celle.value = nyklasse
            antall += 1

    if antall:
        arbeidsbok.save(FILE)  # lagrer én gang til slutt
    return antall


def legg_til(arbeidsbok: Workbook, ark: Worksheet, mail: str, klasse: str, rolle: str) -> None:
    _sjekk_rolle(rolle)
    tabell, min_kol, min_rad, maks_kol, maks_rad = hent_tabell(ark)

    # finn siste rad som faktisk har data (ikke bare stol på ref)
    siste = min_rad
    for r in range(maks_rad, min_rad, -1):
        if ark.cell(row=r, column=min_kol).value not in (None, ""):
            siste = r
            break

    ny_rad = siste + 1
    for kol, verdi in enumerate([mail, klasse, rolle], start=min_kol):
        ark.cell(row=ny_rad, column=kol, value=verdi)

    _sett_ref(tabell, min_kol, min_rad, maks_kol, max(maks_rad, ny_rad))
    arbeidsbok.save(FILE)


def slett(arbeidsbok: Workbook, ark: Worksheet, rad: int) -> None:
    tabell, min_kol, min_rad, maks_kol, maks_rad = hent_tabell(ark)

    if not (min_rad < rad <= maks_rad):
        raise ValueError(f"Rad {rad} er ikke en datarad i tabellen ({min_rad + 1}-{maks_rad}).")
    if maks_rad - min_rad <= 1:
        raise ValueError("Tabellen må ha minst én datarad, så den siste raden kan ikke slettes.")

    ark.delete_rows(rad)
    _sett_ref(tabell, min_kol, min_rad, maks_kol, maks_rad - 1)
    arbeidsbok.save(FILE)


# ---------------------------------------------------------------------------
# Dialog med flere skrivefelt
# ---------------------------------------------------------------------------

class Skjema(tk.Toplevel):
    """Liten dialog med ett tekstfelt per label. .resultat er en liste, eller None hvis avbrutt."""

    def __init__(
        self,
        forelder,
        tittel: str,
        felter: list[str],
        knapp: str = "OK",
        startverdier: list[str] | None = None,
        valg: dict[int, list[str]] | None = None,
    ):
        super().__init__(forelder)
        self.title(tittel)
        self.transient(forelder)
        self.resizable(False, False)
        self.resultat = None
        self.entries: list[tk.Widget] = []
        valg = valg or {}

        ramme = ttk.Frame(self, padding=14)
        ramme.pack()

        for i, tekst in enumerate(felter):
            ttk.Label(ramme, text=tekst).grid(row=i, column=0, sticky="w", pady=4, padx=(0, 12))
            if i in valg:
                # nedtrekksliste: brukeren kan bare velge blant de gyldige verdiene
                entry = ttk.Combobox(ramme, width=29, state="readonly", values=valg[i])
                entry.set(startverdier[i] if startverdier and startverdier[i] in valg[i] else valg[i][0])
            else:
                entry = ttk.Entry(ramme, width=32)
                if startverdier:
                    entry.insert(0, startverdier[i])
            entry.grid(row=i, column=1, pady=4)
            self.entries.append(entry)

        knapper = ttk.Frame(ramme)
        knapper.grid(row=len(felter), column=0, columnspan=2, sticky="e", pady=(12, 0))
        ttk.Button(knapper, text="Avbryt", command=self.destroy).pack(side="right")
        ttk.Button(knapper, text=knapp, command=self._ok).pack(side="right", padx=6)

        self.bind("<Return>", lambda _: self._ok())
        self.bind("<Escape>", lambda _: self.destroy())

        self.entries[0].focus_set()
        if isinstance(self.entries[0], ttk.Entry):
            self.entries[0].select_range(0, "end")
        self.wait_visibility()
        self.grab_set()
        self.wait_window()

    def _ok(self) -> None:
        self.resultat = [e.get().strip() for e in self.entries]
        self.destroy()


# ---------------------------------------------------------------------------
# Hovedvindu
# ---------------------------------------------------------------------------

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Mailliste")
        self.geometry("820x520")
        self.minsize(560, 320)

        try:
            self.last_fil()
        except Exception as feil:
            messagebox.showerror("Kunne ikke åpne filen", f"{FILE}\n\n{feil}")
            self.destroy()
            raise SystemExit

        self.bygg_ui()
        self.oppdater()

    # --- filhåndtering ----------------------------------------------------

    def last_fil(self) -> None:
        self.arbeidsbok: Workbook = load_workbook(FILE)
        self.ark: Worksheet = self.arbeidsbok.active

    def kjor(self, handling, *args):
        """Kjører en Excel-funksjon og viser feil. Returnerer (ok, resultat)."""
        try:
            return True, handling(self.arbeidsbok, self.ark, *args)
        except ValueError as feil:
            messagebox.showwarning("Kan ikke gjøre det", str(feil))
        except PermissionError:
            messagebox.showerror(
                "Kunne ikke lagre",
                "Filen er åpen i et annet program. Lukk den i Excel og prøv igjen.",
            )
            self.last_fil()  # forkast endringer som ikke ble lagret
            self.oppdater()
        except Exception as feil:
            messagebox.showerror("Noe gikk galt", str(feil))
            self.last_fil()
            self.oppdater()
        return False, None

    # --- grensesnitt --------------------------------------------------------

    def bygg_ui(self) -> None:
        _, min_kol, min_rad, maks_kol, _ = hent_tabell(self.ark)
        self.overskrifter = [
            str(self.ark.cell(row=min_rad, column=k).value or "")
            for k in range(min_kol, maks_kol + 1)
        ]

        knapper = ttk.Frame(self, padding=(10, 10, 10, 4))
        knapper.pack(fill="x")
        ttk.Button(knapper, text="Legg til", command=self.legg_til_klikk).pack(side="left")
        ttk.Button(knapper, text="Endre valgt rad", command=self.endre_rad_klikk).pack(side="left", padx=6)
        ttk.Button(knapper, text="Slett valgt rad", command=self.slett_klikk).pack(side="left")
        ttk.Button(knapper, text="Endre klasse", command=self.endre_alle_klikk).pack(side="left", padx=6)

        midt = ttk.Frame(self, padding=(10, 0))
        midt.pack(fill="both", expand=True)

        self.tre = ttk.Treeview(
            midt, columns=list(range(len(self.overskrifter))), show="headings", selectmode="browse"
        )
        for i, navn in enumerate(self.overskrifter):
            self.tre.heading(i, text=navn)
            self.tre.column(i, width=220, anchor="w")

        rull = ttk.Scrollbar(midt, orient="vertical", command=self.tre.yview)
        self.tre.configure(yscrollcommand=rull.set)
        self.tre.pack(side="left", fill="both", expand=True)
        rull.pack(side="right", fill="y")

        self.tre.bind("<Double-1>", self.dobbeltklikk)

        self.status = tk.StringVar(value="Dobbeltklikk på en celle for å endre den.")
        ttk.Label(self, textvariable=self.status, padding=(10, 6)).pack(fill="x")

    def oppdater(self) -> None:
        """Tegner tabellen på nytt fra arket."""
        self.tre.delete(*self.tre.get_children())
        _, min_kol, min_rad, maks_kol, maks_rad = hent_tabell(self.ark)

        for rad in self.ark.iter_rows(
            min_row=min_rad + 1, max_row=maks_rad, min_col=min_kol, max_col=maks_kol
        ):
            verdier = ["" if c.value is None else c.value for c in rad]
            self.tre.insert("", "end", iid=str(rad[0].row), values=verdier)  # iid = radnummer i arket

    # --- handlinger -----------------------------------------------------------

    def dobbeltklikk(self, event) -> None:
        if self.tre.identify_region(event.x, event.y) != "cell":
            return

        rad = int(self.tre.identify_row(event.y))
        kol_idx = int(self.tre.identify_column(event.x)[1:]) - 1  # "#2" -> 1

        _, min_kol, *_ = hent_tabell(self.ark)
        kolonne = min_kol + kol_idx

        gammel = self.ark.cell(row=rad, column=kolonne).value
        gammel_tekst = "" if gammel is None else str(gammel)

        if kol_idx == ROLLE_KOLONNE:
            svar = Skjema(
                self,
                "Endre celle",
                [f"Rolle (rad {rad})"],
                knapp="Lagre",
                startverdier=[gammel_tekst],
                valg={0: list(GYLDIGE_ROLLER)},
            ).resultat
            ny = svar[0] if svar else None
        else:
            ny = simpledialog.askstring(
                "Endre celle",
                f"Ny verdi for «{self.overskrifter[kol_idx]}» (rad {rad}):",
                initialvalue=gammel_tekst,
                parent=self,
            )
        if ny is None or ny == gammel_tekst:
            return

        ok, _ = self.kjor(endre_en, rad, kolonne, ny)
        if ok:
            self.oppdater()
            self.status.set(f"Endret rad {rad}: «{gammel_tekst}» → «{ny}»")

    def endre_rad_klikk(self) -> None:
        valgt = self.tre.selection()
        if not valgt:
            messagebox.showinfo("Endre rad", "Klikk på en rad først.")
            return

        rad = int(valgt[0])
        _, min_kol, _, maks_kol, _ = hent_tabell(self.ark)

        # les fra arket (ikke fra visningen) så verdiene er nøyaktig som i filen
        gamle = []
        for k in range(min_kol, maks_kol + 1):
            v = self.ark.cell(row=rad, column=k).value
            gamle.append("" if v is None else str(v))

        svar = Skjema(
            self,
            f"Endre rad {rad}",
            self.overskrifter,
            knapp="Lagre",
            startverdier=gamle,
            valg={ROLLE_KOLONNE: list(GYLDIGE_ROLLER)},
        ).resultat
        if svar is None:
            return
        if svar == gamle:
            self.status.set("Ingen endringer.")
            return
        if not svar[0]:
            messagebox.showwarning("Mangler verdi", f"«{self.overskrifter[0]}» kan ikke være tom.")
            return

        ok, _ = self.kjor(endre_rad, rad, svar)
        if ok:
            self.oppdater()
            self.tre.selection_set(str(rad))  # behold valgt rad
            self.status.set(f"Endret rad {rad}.")

    def endre_alle_klikk(self) -> None:
        svar = Skjema(
            self,
            "Endre alle",
            ["Hvilken klasse vil du endre?", "Hva heter den nye klassen?"],
            knapp="Endre",
        ).resultat
        if svar is None:
            return

        gammel, ny = svar
        if not gammel or not ny:
            messagebox.showwarning("Mangler verdi", "Fyll ut begge feltene.")
            return

        ok, antall = self.kjor(endre_alle, gammel, ny)
        if ok:
            self.oppdater()
            self.status.set(f"Endret {antall} rad(er) fra «{gammel}» til «{ny}».")

    def legg_til_klikk(self) -> None:
        svar = Skjema(
            self,
            "Legg til rad",
            self.overskrifter,
            knapp="Legg til",
            valg={ROLLE_KOLONNE: list(GYLDIGE_ROLLER)},
        ).resultat
        if svar is None:
            return

        if not svar[0]:
            messagebox.showwarning("Mangler verdi", f"«{self.overskrifter[0]}» kan ikke være tom.")
            return

        ok, _ = self.kjor(legg_til, *svar)
        if ok:
            self.oppdater()
            self.status.set(f"La til «{svar[0]}».")

    def slett_klikk(self) -> None:
        valgt = self.tre.selection()
        if not valgt:
            messagebox.showinfo("Slett", "Klikk på en rad først.")
            return

        rad = int(valgt[0])
        tekst = " | ".join(str(v) for v in self.tre.item(valgt[0], "values"))
        if not messagebox.askyesno("Slett rad", f"Vil du slette denne raden?\n\n{tekst}"):
            return

        ok, _ = self.kjor(slett, rad)
        if ok:
            self.oppdater()
            self.status.set(f"Slettet rad {rad}.")


if __name__ == "__main__":
    App().mainloop()